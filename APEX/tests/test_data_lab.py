"""Local dummy-data intake: exact schema, row exclusion, privacy, and comparison."""
import csv
import io

import pytest
from openpyxl import load_workbook

from data_lab import (TEMPLATE_COLUMNS, batch_compare, synthetic_sample_csv, synthetic_sample_xlsx,
                      template_csv, template_xlsx, validate_csv, validate_xlsx, validated_history)


def _edit_rows(raw: bytes, edit):
    reader = csv.DictReader(io.StringIO(raw.decode("utf-8")))
    rows = list(reader)
    edit(rows)
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=reader.fieldnames)
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode("utf-8")


def test_synthetic_template_validates_and_compares_without_pan(bank_b):
    assert tuple(next(csv.reader(io.StringIO(template_csv().decode("utf-8"))))) == TEMPLATE_COLUMNS
    intake = validate_csv(synthetic_sample_csv(bank_b, 12))
    assert intake.total_rows == len(intake.valid) == 12
    assert not intake.errors
    comparison = batch_compare(intake, bank_b)
    assert len(comparison) == 12
    assert {"bank_decision", "apex_classification", "agreement", "fail_checks"} <= set(comparison)
    assert "pan" not in comparison.columns


def test_errors_exclude_exact_rows_without_echoing_sensitive_values(bank_b):
    raw = synthetic_sample_csv(bank_b, 20)

    def corrupt(rows):
        rows[1]["net_monthly_income"] = "not-a-number"
        rows[3]["decision_reason"] = "1234 5678 9012"

    intake = validate_csv(_edit_rows(raw, corrupt))
    assert intake.total_rows == 20
    assert len(intake.valid) == 18
    assert {e["row"] for e in intake.errors} == {3, 5}
    assert "1234 5678 9012" not in str(intake.errors)
    assert "not-a-number" not in str(intake.errors)


def test_sensitive_or_unknown_columns_are_rejected(bank_b):
    raw = synthetic_sample_csv(bank_b, 2).decode("utf-8")
    with pytest.raises(ValueError, match="Sensitive columns"):
        validate_csv(raw.replace("applicant_id,", "customer_name,applicant_id,", 1).encode("utf-8"))
    with pytest.raises(ValueError, match="unknown columns"):
        validate_csv(raw.replace("decision_reason,", "extra,decision_reason,", 1).encode("utf-8"))


def test_explicit_money_units_are_applied(bank_b):
    raw = synthetic_sample_csv(bank_b, 1)
    original = validate_csv(raw).valid[0].applicant.requested_amount

    def to_thousands(rows):
        rows[0]["requested_amount"] = str(float(rows[0]["requested_amount"]) / 1000)
        rows[0]["monthly_income_observations"] = "[1, 2]"

    changed = _edit_rows(raw, to_thousands)
    parsed = validate_csv(changed, unit_multiplier=1000)
    assert len(parsed.valid) == 1
    # Conversion applies to every monetary column, not only the requested amount.
    assert parsed.valid[0].applicant.requested_amount == original
    assert parsed.valid[0].applicant.monthly_income_observations == [1000.0, 2000.0]


def test_pan_is_hashed_in_memory_and_not_in_comparison(bank_b):
    raw_pan = "ABCDE1234F"
    sample = synthetic_sample_csv(bank_b, 1)

    def add_pan(rows):
        rows[0]["pan"] = raw_pan
        rows[0]["pan_consent"] = "True"

    intake = validate_csv(_edit_rows(sample, add_pan))
    assert len(intake.valid) == 1
    assert intake.valid[0].pan_hmac and intake.valid[0].pan_hmac != raw_pan
    assert raw_pan not in batch_compare(intake, bank_b).to_csv(index=False)
    assert raw_pan not in validated_history(intake).to_csv(index=False)
    assert "pan" not in validated_history(intake).columns


def test_five_percent_planted_errors_are_excluded(bank_b):
    sample = synthetic_sample_csv(bank_b, 5000)

    def plant(rows):
        for row in rows[::20]:
            row["decision"] = "invalid"

    intake = validate_csv(_edit_rows(sample, plant))
    assert intake.total_rows == 5000
    assert len(intake.valid) == 4750
    assert len({e["row"] for e in intake.errors}) == 250


def test_excel_template_and_sample_use_the_same_validated_contract(bank_b):
    template = load_workbook(io.BytesIO(template_xlsx()), read_only=True)
    assert template.sheetnames == ["Applicants"]
    assert tuple(c.value for c in template.active[1]) == TEMPLATE_COLUMNS
    template.close()
    intake = validate_xlsx(synthetic_sample_xlsx(bank_b, 15))
    assert intake.total_rows == len(intake.valid) == 15
    assert not intake.errors


def test_excel_formulas_are_rejected(bank_b):
    book = load_workbook(io.BytesIO(synthetic_sample_xlsx(bank_b, 1)))
    book.active["A2"] = '=HYPERLINK("https://example.test", "row")'
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(ValueError, match="formulas are not accepted"):
        validate_xlsx(out.getvalue())


def test_birth_date_after_decision_date_excludes_row(bank_b):
    sample = synthetic_sample_csv(bank_b, 1)
    changed = _edit_rows(sample, lambda rows: rows[0].update(date_of_birth="2030-01-01"))
    intake = validate_csv(changed)
    assert len(intake.valid) == 0
    assert intake.errors[0]["row"] == 2


def test_wrong_policy_and_duplicate_application_are_excluded(bank_b):
    sample = synthetic_sample_csv(bank_b, 3)

    def change(rows):
        rows[1]["policy_version"] = "another-version"
        rows[2]["applicant_id"] = rows[0]["applicant_id"]

    intake = validate_csv(_edit_rows(sample, change), cfg=bank_b)
    assert len(intake.valid) == 1
    assert {e["row"] for e in intake.errors} == {3, 4}
    assert {e["field"] for e in intake.errors} == {"policy_version", "applicant_id"}
