"""Local-only intake and comparison for synthetic or bank-supplied dummy CSVs.

No uploaded bytes, applicant rows or PAN values are written to disk here.
The comparison describes agreement with historical decisions, not accuracy.
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from engine import assess
from age_math import completed_years
from kyc import pan_hash
from models import ApplicantInput, BankConfig, CheckStatus, Classification
from synthetic_history import generate

META_COLUMNS = ("decision", "decision_date", "decision_reason", "product", "policy_version",
                "branch_hash", "officer_hash", "dpd_6m_on_book", "dpd_12m_on_book", "dpd_24m_on_book")
INPUT_COLUMNS = tuple(ApplicantInput.model_fields)
TEMPLATE_COLUMNS = INPUT_COLUMNS + META_COLUMNS
_REQUIRED = {"applicant_id", "requested_amount", "tenure_months", "annual_interest_rate",
             "decision", "decision_date"}
_MONEY_FIELDS = {"requested_amount", "net_monthly_income", "verified_monthly_income",
                 "existing_monthly_emis", "other_recurring_obligations", "credit_card_limit",
                 "credit_card_balance", "avg_monthly_balance", "other_upfront_charges"}
_SENSITIVE_HEADER = re.compile(r"(^|_)(name|first_name|last_name|phone|mobile|aadhaar|aadhar)(_|$)", re.I)
_AADHAAR_LIKE = re.compile(r"(?<!\d)\d{4}[ -]?\d{4}[ -]?\d{4}(?!\d)")


@dataclass(frozen=True)
class ValidRow:
    source_row: int
    applicant: ApplicantInput
    decision: int
    decision_date: date
    pan_hmac: str | None


@dataclass(frozen=True)
class IntakeResult:
    valid: list[ValidRow]
    errors: list[dict[str, object]]
    total_rows: int


def template_csv() -> bytes:
    out = io.StringIO()
    csv.writer(out).writerow(TEMPLATE_COLUMNS)
    return out.getvalue().encode("utf-8")


def _xlsx_from_rows(rows) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Applicants"
    sheet.append(TEMPLATE_COLUMNS)
    for row in rows:
        sheet.append([row.get(name, "") for name in TEMPLATE_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="173A70")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(28, max(16, len(str(column[0].value)) + 2))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def template_xlsx() -> bytes:
    return _xlsx_from_rows(())


def synthetic_sample_xlsx(cfg: BankConfig, n: int = 200) -> bytes:
    rows = csv.DictReader(io.StringIO(synthetic_sample_csv(cfg, n).decode("utf-8")))
    return _xlsx_from_rows(rows)


def validate_xlsx(data: bytes, unit_multiplier: int = 1, cfg: BankConfig | None = None) -> IntakeResult:
    """Convert the Applicants worksheet to the exact CSV contract before validation."""
    if len(data) > 25_000_000:
        raise ValueError("XLSX exceeds the 25 MB local intake limit")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("XLSX could not be opened") from exc
    try:
        if book.sheetnames != ["Applicants"]:
            raise ValueError("XLSX must contain one sheet named Applicants")
        sheet = book["Applicants"]
        if sheet.max_row > 100_001 or sheet.max_column > len(TEMPLATE_COLUMNS):
            raise ValueError("XLSX exceeds the 100,000-row template limit")
        out = io.StringIO()
        writer = csv.writer(out)
        for row in sheet.iter_rows(values_only=False):
            values = []
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError("XLSX formulas are not accepted; paste values instead")
                value = cell.value
                if isinstance(value, (date, datetime)):
                    value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                values.append("" if value is None else str(value))
            writer.writerow(values)
        return validate_csv(out.getvalue().encode("utf-8"), unit_multiplier, cfg)
    finally:
        book.close()


def synthetic_sample_csv(cfg: BankConfig, n: int = 200) -> bytes:
    """A repeatable file with every template column and no real identifiers."""
    frame = generate(cfg, n=n, seed=17)
    for name in TEMPLATE_COLUMNS:
        if name not in frame:
            frame[name] = ""
    frame["product"] = cfg.meta.product
    frame["policy_version"] = cfg.meta.policy_version
    frame["branch_hash"] = "SYNTHETIC-BRANCH"
    frame["officer_hash"] = "SYNTHETIC-OFFICER"
    return frame.loc[:, TEMPLATE_COLUMNS].to_csv(index=False).encode("utf-8")


def _input_value(name: str, raw: str):
    if raw == "":
        return None
    if name in {"documents", "monthly_income_observations"}:
        return json.loads(raw)
    return raw


def validate_csv(data: bytes, unit_multiplier: int = 1, cfg: BankConfig | None = None) -> IntakeResult:
    """Validate without repairing rows; returned errors never contain input values."""
    if unit_multiplier not in (1, 1_000, 100_000):
        raise ValueError("Unit must be rupees, thousands of rupees or lakhs")
    if len(data) > 25_000_000:
        raise ValueError("CSV exceeds the 25 MB local intake limit")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV must be UTF-8 encoded") from exc
    reader = csv.DictReader(io.StringIO(text, newline=""))
    headers = reader.fieldnames or []
    if len(headers) != len(set(headers)):
        raise ValueError("CSV has duplicate column names")
    sensitive = [h for h in headers if _SENSITIVE_HEADER.search(h)]
    if sensitive:
        raise ValueError("Sensitive columns are not accepted: " + ", ".join(sensitive))
    unknown = set(headers) - set(TEMPLATE_COLUMNS)
    missing = set(TEMPLATE_COLUMNS) - set(headers)
    if unknown or missing:
        parts = []
        if missing:
            parts.append("missing columns: " + ", ".join(sorted(missing)))
        if unknown:
            parts.append("unknown columns: " + ", ".join(sorted(unknown)))
        raise ValueError("; ".join(parts))

    valid: list[ValidRow] = []
    errors: list[dict[str, object]] = []
    total = 0
    seen_ids: set[str] = set()
    for number, row in enumerate(reader, start=2):
        total += 1
        if None in row or any(value is None for value in row.values()):
            errors.append({"row": number, "field": "CSV", "issue": "column count differs from header"})
            continue
        unsafe = next((name for name, value in row.items() if name != "pan" and
                       _AADHAAR_LIKE.search(value)), None)
        if unsafe:
            errors.append({"row": number, "field": unsafe, "issue": "possible 12-digit identifier"})
            continue
        missing_values = [name for name in _REQUIRED if not row[name].strip()]
        if missing_values:
            errors.extend({"row": number, "field": name, "issue": "required value missing"}
                          for name in sorted(missing_values))
            continue
        if cfg is not None:
            mismatch = (["product"] if row["product"] != cfg.meta.product else []) + (
                ["policy_version"] if row["policy_version"] != cfg.meta.policy_version else [])
            if mismatch:
                errors.extend({"row": number, "field": name,
                               "issue": "does not match the selected bank policy"} for name in mismatch)
                continue
        try:
            decision = int(row["decision"])
            if decision not in (0, 1) or row["decision"] not in ("0", "1"):
                raise ValueError("decision must be 0 or 1")
            decided_on = date.fromisoformat(row["decision_date"])
            payload = {name: _input_value(name, row[name]) for name in INPUT_COLUMNS
                       if row[name] != ""}
            if unit_multiplier != 1:
                for name in _MONEY_FIELDS & payload.keys():
                    payload[name] = float(payload[name]) * unit_multiplier
                if isinstance(payload.get("monthly_income_observations"), list):
                    payload["monthly_income_observations"] = [
                        float(value) * unit_multiplier for value in payload["monthly_income_observations"]]
            applicant = ApplicantInput.model_validate(payload)
            if applicant.date_of_birth is not None:
                completed_years(applicant.date_of_birth, decided_on)
        except ValidationError as exc:
            errors.extend({"row": number, "field": ".".join(map(str, e["loc"])) or "input",
                           "issue": e["type"]} for e in exc.errors())
            continue
        except (ValueError, TypeError, json.JSONDecodeError):
            errors.append({"row": number, "field": "input or decision", "issue": "invalid type or format"})
            continue
        if applicant.applicant_id in seen_ids:
            errors.append({"row": number, "field": "applicant_id", "issue": "duplicate application ID"})
            continue
        seen_ids.add(applicant.applicant_id)
        valid.append(ValidRow(number, applicant, decision, decided_on, pan_hash(applicant.pan)))
    return IntakeResult(valid, errors, total)


def batch_compare(intake: IntakeResult, cfg: BankConfig) -> pd.DataFrame:
    """Return a PAN-free, local summary for each valid applicant."""
    rows = []
    for item in intake.valid:
        result = assess(item.applicant, cfg,
                        timestamp=item.decision_date.isoformat() + "T00:00:00+00:00", use_model=False)
        status = result.classification
        agreement = ("same direction" if (status == Classification.LIKELY and item.decision == 1) or
                     (status == Classification.UNLIKELY and item.decision == 0) else
                     "different direction" if status in {Classification.LIKELY, Classification.UNLIKELY} else
                     "manual review / not comparable")
        fails = [c.label for c in result.checks if c.status == CheckStatus.FAIL]
        reviews = [c.label for c in result.checks if c.status in {CheckStatus.REVIEW, CheckStatus.DEVIATION}]
        rows.append({"source_row": item.source_row, "applicant_id": item.applicant.applicant_id,
                     "bank_decision": "approved" if item.decision else "declined",
                     "apex_classification": status.value, "agreement": agreement,
                     "fail_checks": "; ".join(fails), "review_checks": "; ".join(reviews),
                     "adjusted_income": result.income.adjusted_income,
                     "foir": result.metrics["foir"].value,
                     "missing_checks": sum(c.missing_info for c in result.checks)})
    return pd.DataFrame(rows)


def validated_history(intake: IntakeResult) -> pd.DataFrame:
    """PAN-free history for the existing Model Lab, from validated rows only."""
    rows = []
    for item in intake.valid:
        row = item.applicant.model_dump(mode="json", exclude={"pan"})
        row.update(decision=item.decision, decision_date=item.decision_date.isoformat())
        rows.append(row)
    return pd.DataFrame(rows)
